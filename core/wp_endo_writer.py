"""
wp_endo_writer.py — Writes EndoRecord lists into a WP12345 EcoTEA Endo template.

The WP12345 'Sheet 1' has 9 header/metadata rows (rows 1-9 in openpyxl 1-based).
Data rows start at row 10. This writer preserves all header rows and writes
EndoRecord objects from row 10 onwards.
"""

import os
import shutil

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from core.base_model import EndoRecord, EMPTY


# ── Column mapping: EndoRecord field → openpyxl column index (1-based) ───────
ENDO_COL_ORDER = [
    'wp6_title',            # col 1   (A)
    'data_owner',           # col 2   (B)
    'data_provider',        # col 3   (C)
    'data_source',          # col 4   (D)
    'data_source_desc',     # col 5   (E)
    'data_user',            # col 6   (F)
    'usage_purpose',        # col 7   (G)
    'process_code',         # col 8   (H)
    'description',          # col 9   (I)
    'geography',            # col 10  (J)
    'year',                 # col 11  (K)
    'start_year',           # col 12  (L)
    'lifetime',             # col 13  (M)
    'grade',                # col 14  (N)
    'ef',                   # col 15  (O)
    'ef_unit',              # col 16  (P)
    'currency',             # col 17  (Q)
    'capex',                # col 18  (R)
    'capex_unit',           # col 19  (S)
    'fixed_opex',           # col 20  (T)
    'fixed_opex_unit',      # col 21  (U)
    'variable_opex',        # col 22  (V)
    'variable_opex_unit',   # col 23  (W)
    'currency_opex',        # col 24  (X)
    'opex',                 # col 25  (Y)
    'tax_cost',             # col 26  (Z)
    'sub_cost',             # col 27  (AA)
    'fixed_annuity',        # col 28  (AB)
    'capex_annuity',        # col 29  (AC)
    'start_cost',           # col 30  (AD)
    'efficiency',           # col 31  (AE)
    'efficiency_unit',      # col 32  (AF)
    'tech_efficiency',      # col 33  (AG)
    'tech_efficiency_unit', # col 34  (AH)
    'commodity_share',      # col 35  (AI)
    'commodity',            # col 36  (AJ)
    'commodity_demand',     # col 37  (AK)
    'commodity_cost_base',  # col 38  (AL)
    'commodity_cost_low',   # col 39  (AM)
    'maintenance_rate',     # col 40  (AN)
    'maintenance_schedule', # col 41  (AO)
    'mean_time_to_repair',  # col 42  (AP)
    'interpolation_rule',   # col 43  (AQ)
    'afa',                  # col 44  (AR)
    'heat_rate',            # col 45  (AS)
    'forced_outage_rate',   # col 46  (AT)
    'max_rating_factor',    # col 47  (AU)
    'ramp_rate',            # col 48  (AV)
    'min_stable_load',      # col 49  (AW)
    'battery_capacity',     # col 50  (AX)
    'max_range',            # col 51  (AY)
    'max_ac_rate',          # col 52  (AZ)
    'max_dc_rate',          # col 53  (BA)
    'fuel_type',            # col 54  (BB)
    'ccs_rate',             # col 55  (BC)
    'temperature_class',    # col 56  (BD)
    'capacity',             # col 57  (BE)
    'capacity_type',        # col 58  (BF)
    'extra1',               # col 59  (BG)
    'extra2',               # col 60  (BH)
    'extra3',               # col 61  (BI)
]

# WP12345 'Sheet 1' header rows: rows 1-9 are preserved; data starts at row 10
ENDO_DATA_START_ROW = 10
ENDO_SHEET_NAME = 'Sheet 1'


def write_endo_output(records: list[EndoRecord],
                      template_path: str,
                      output_path: str) -> str:
    """
    Write EndoRecord rows into the WP12345 Endo template.

    Preserves all header rows (1-9). Clears any existing data rows (10+),
    then writes the supplied records from row 10 onwards.

    Args:
        records:       List of EndoRecord objects.
        template_path: Path to the WP12345 EcoTEA Endo .xlsx template.
        output_path:   Destination path for the filled output .xlsx.

    Returns:
        output_path as string.
    """
    # Copy file data only (not permissions) so the output is always writable
    shutil.copyfile(template_path, output_path)
    os.chmod(output_path, 0o644)

    wb = openpyxl.load_workbook(output_path)
    ws = wb[ENDO_SHEET_NAME]

    # Clear any existing data rows from the template
    max_row = ws.max_row
    if max_row >= ENDO_DATA_START_ROW:
        for row in ws.iter_rows(min_row=ENDO_DATA_START_ROW, max_row=max_row):
            for cell in row:
                cell.value = None

    # Write records starting at data row
    for row_offset, rec in enumerate(records):
        row_idx = ENDO_DATA_START_ROW + row_offset
        for col_idx, field_name in enumerate(ENDO_COL_ORDER, start=1):
            val = getattr(rec, field_name, EMPTY)

            # EMPTY (None) → leave cell blank; any other value is written as-is
            if val is EMPTY or val is None:
                val = None
            elif isinstance(val, float) and str(val) == 'nan':
                val = None

            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.alignment = Alignment(horizontal='left', vertical='center',
                                       wrap_text=False)

    wb.save(output_path)
    return output_path
